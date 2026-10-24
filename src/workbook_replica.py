"""
Python replica of the ORIGINAL Excel workbook logic (MINOR_PROJECT_DONE_BY_KRUSHNA.xlsm),
reproduced formula-by-formula so that its results can be compared with the IS 10262-based
calculator. The workbook file itself is never modified.

Also contains `read_workbook_inputs()` (openpyxl) and the AUDIT_FINDINGS list shown in the app.
"""
from __future__ import annotations

from typing import Dict, List

# TABLES sheet as stored in the workbook -----------------------------------------------
WB_S = {"M10": 3.5, "M15": 3.5, "M20": 4, "M25": 4, "M30": 4, "M35": 5, "M40": 5, "M45": 5,
        "M50": 5, "M55": 5, "M60": 5, "M65": 5, "M70": 5, "M75": 5, "M80": 5}
WB_X = {"M10": 5, "M15": 5, "M20": 5, "M25": 6, "M30": 6, "M35": 7, "M40": 7, "M45": 8,
        "M50": 8, "M55": 9, "M60": 9, "M65": 10, "M70": 10, "M75": 11, "M80": 11}
WB_WC = {"MILD": 0.55, "MODERATE": 0.5, "SEVERE": 0.45}
WB_WATER = {50: 186, 75: 191, 100: 197, 120: 202}
WB_SLUMP_PCT = {50: 0, 75: 3, 100: 6, 120: 9}
WB_SG = {"NORMAL": 2.7, "LECA": 1.2, "PUMIC": 1.5}
WB_WA = {"NORMAL": 2, "LECA": 10, "PUMIC": 8}
WB_AIR = {10: 3, 20: 2, 40: 1}


def run_workbook_logic(grade="M70", method="S", exposure="SEVERE", slump=120, agg_size=10,
                       agg_type="NORMAL", condition="SSD", fly_ash_pct=10, plasticizer_pct=1,
                       sg_cement=3.15, fine_pct=50, coarse_pct=50) -> Dict[str, float]:
    """Mirror of the CALCULATIONS and OUTPUT sheets (cell references in comments)."""
    fck = float(grade[1:4])                                   # C2 = VALUE(MID(B2,2,3))
    S, X = WB_S[grade], WB_X[grade]                           # C3, C4
    target = fck + 1.65 * S if method == "S" else fck + X     # C5
    if fck <= 25:                                             # C6
        wc = WB_WC[exposure]
    elif fck <= 30:
        wc = 0.45
    elif fck <= 40:
        wc = 0.4
    elif fck <= 50:
        wc = 0.38
    else:
        wc = 0.35
    base = WB_WATER[slump]                                    # C7 (already includes slump!)
    corr = WB_SLUMP_PCT[slump] / 100                          # C8
    corrected = base * (1 + corr)                             # C9  <- slump counted twice
    after_pl = corrected * (1 - plasticizer_pct / 100)        # C10 <- dose treated as water reduction
    cement = after_pl / wc                                    # C11
    flyash = cement * fly_ash_pct / 100                       # C12
    net_cement = cement - flyash                              # C13
    sg_agg = WB_SG[agg_type]                                  # C14
    wa = WB_WA[agg_type]                                      # C15
    air = WB_AIR[agg_size] / 100                              # C16
    vol = 1 - (net_cement / (sg_cement * 1000) + after_pl / 1000 + air)   # C17 <- no FA / admixture volume
    total_agg = vol * sg_agg * 1000                           # C18
    fine = total_agg * fine_pct / 100                         # C19
    coarse = total_agg * coarse_pct / 100                     # C20
    extra = total_agg * wa / 100                              # C21
    final_water = after_pl + extra if condition == "Dry" else after_pl   # C22
    plast_qty = net_cement * plasticizer_pct / 100            # C23
    # OUTPUT sheet status bug: B12 = VALUE(MID(B2,2,3))*10 where B2 is the NUMERIC target strength
    mid = str(target)[1:4]
    try:
        status_threshold = float(mid) * 10
    except ValueError:
        status_threshold = float("nan")
    density = net_cement + flyash + final_water + fine + coarse   # B17
    return {
        "target_strength": target, "wc": wc, "water_table": base, "water_after_plasticizer": after_pl,
        "cement_net": net_cement, "fly_ash": flyash, "total_binder": cement,
        "fine_agg": fine, "coarse_agg": coarse, "final_water": final_water,
        "plasticizer": plast_qty, "aggregate_volume_m3": vol, "density_check": density,
        "status_threshold_used_by_B12": status_threshold,
    }


AUDIT_FINDINGS: List[Dict[str, str]] = [
    {"id": "W1", "sheet": "CALCULATIONS C3/C4 + TABLES",
     "issue": "Standard deviation S and factor X tables differ from IS 10262:2019 Tables 1-2.",
     "detail": "Workbook has S=5 for M65-M80 (IS: 6.0) and X=7 for M35/M40 (IS: 6.5), X=8..11 for "
               "M45-M80 (IS: 6.5 up to M60, 8.0 for M65+), X=5 for M20 (IS: 5.5), X=6 for M25 (5.5).",
     "severity": "High"},
    {"id": "W2", "sheet": "CALCULATIONS C5",
     "issue": "Target strength uses fck+1.65S OR fck+X chosen by a flag ('S'), not the larger of both.",
     "detail": "IS 10262 cl. 4.2 requires f'ck = the HIGHER of the two.", "severity": "High"},
    {"id": "W3", "sheet": "CALCULATIONS C6",
     "issue": "w/c is a hard-coded step function of grade, not read from Fig. 1 / Table 8 and not "
              "linked to the target strength or cement grade.",
     "detail": "e.g. every grade above M50 gets 0.35 regardless of M55 vs M80.", "severity": "High"},
    {"id": "W4", "sheet": "CALCULATIONS C7-C9",
     "issue": "Slump is applied twice.",
     "detail": "TABLES J:K already lists 186/191/197/202 kg (which include the slump increase); "
               "C8 then multiplies by a further 0/3/6/9 %. For 120 mm: 202 x 1.09 = 220 kg instead of ~202.",
     "severity": "High"},
    {"id": "W5", "sheet": "CALCULATIONS C7",
     "issue": "Water content ignores nominal maximum aggregate size and aggregate shape.",
     "detail": "IS 10262 Table 4 / Table 7 depend on msa (10/20/40 mm); shape reductions of 10-20 kg apply.",
     "severity": "Medium"},
    {"id": "W6", "sheet": "CALCULATIONS C10",
     "issue": "Plasticiser dose % is subtracted directly as water reduction.",
     "detail": "Dose and water reduction are different quantities (Annex G: 0.5-1.5 % dose gives "
               "15-30 %+ reduction, found by trials). 1 % dose -> 1 % reduction is unsupported.",
     "severity": "High"},
    {"id": "W7", "sheet": "CALCULATIONS C11-C13",
     "issue": "Fly ash is carved out of the cement mass, and no 10 % binder increase / durability "
              "minimum / maximum-cement check is made.",
     "detail": "Acceptable only if intended; but Fly-ash and admixture volumes are then ignored (W8).",
     "severity": "Medium"},
    {"id": "W8", "sheet": "CALCULATIONS C17",
     "issue": "Absolute-volume equation omits fly ash and admixture volumes.",
     "detail": "Aggregate volume is over-estimated whenever fly ash or plasticiser is used.",
     "severity": "High"},
    {"id": "W9", "sheet": "CALCULATIONS C14-C20 / INPUT B15:B16",
     "issue": "One SG/absorption per 'aggregate type' and a manual fine:coarse split (e.g. 50:50).",
     "detail": "IS 10262 Tables 5/10 give the coarse fraction from msa, fine-aggregate zone and w/cm "
               "(0.01 per 0.05 w/cm). SG of coarse and fine are separate (2.74 / 2.65 in Annexes).",
     "severity": "High"},
    {"id": "W10", "sheet": "TABLES W:X / INPUT B14",
     "issue": "Air content table (3/2/1 %) is not the IS value (1.5/1.0/0.8 %) and INPUT!B14 (air) is never used.",
     "detail": "C16 reads the TABLES lookup, ignoring the input cell.", "severity": "Medium"},
    {"id": "W11", "sheet": "CALCULATIONS C21-C22",
     "issue": "Extra water for dry aggregate is applied to total aggregate with a single absorption "
              "(2 % for 'NORMAL'); IS examples use 0.5 % (CA) and 1.0 % (FA), and aggregate masses "
              "are not converted to dry masses.",
     "detail": "Per Annex A-11: dry mass = SSD/(1+abs), extra water = difference, per aggregate.",
     "severity": "Medium"},
    {"id": "W12", "sheet": "OUTPUT B12, B14",
     "issue": "Status formula applies VALUE(MID(B2,2,3)) to a NUMERIC target strength.",
     "detail": "MID of 78.25 returns '8.2' -> threshold 82 kg/m3, so the 'cement >= grade x 10' test is "
               "meaningless (and the rule itself is not in IS 10262).", "severity": "High"},
    {"id": "W13", "sheet": "OUTPUT B17:B18",
     "issue": "Density check sums only some components (admixture excluded) against an arbitrary 2300-2600 window.",
     "detail": "Use mass balance + absolute-volume identity; density window is only a plausibility hint.",
     "severity": "Low"},
    {"id": "W14", "sheet": "VBA project",
     "issue": "The .xlsm contains a vbaProject.bin whose module list shows only the default Sheet/ThisWorkbook "
              "stubs; no custom macro source could be extracted here (no oletools available offline).",
     "detail": "VBA functionality could not be inspected in detail - disclosed limitation.", "severity": "Info"},
    {"id": "R1", "sheet": "Project report section 4.1",
     "issue": "Comparison table labels water (191.5, 201, 141 kg) as 'W/C'; the 'Coarse aggregate' of the "
              "developed sheet is 1021.67 for M40/75 mm vs IS 1228; the Excel M40 input sheet shows fly ash 0 % "
              "but the compared IS example B uses 30 %.",
     "detail": "Re-label columns and compare like-for-like inputs before presenting.", "severity": "Medium"},
]


def read_workbook_inputs(path: str) -> Dict[str, object]:
    """Read the INPUT sheet values with openpyxl (formulas of other sheets are not evaluated)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, keep_vba=False)
    ws = wb["INPUT"]
    out = {}
    for r in range(2, ws.max_row + 1):
        k, v = ws.cell(r, 1).value, ws.cell(r, 2).value
        if k:
            out[str(k).strip()] = v
    return out
